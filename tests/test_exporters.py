from scrape_engine.exporters.excel import export_xlsx
from scrape_engine.exporters.json_out import export_json
from scrape_engine.models import EXCEL_HEADERS, Product, Variant, flatten_product
from pathlib import Path


def test_excel_headers_match_sample_order(tmp_path: Path):
    product = Product(
        name="P",
        source_link="https://www.tokopedia.com/a/b",
        variants=[Variant(options=[("Ukuran", "M")], price=1, discount=2, stock=3, sku="S")],
    )
    rows = flatten_product(product)
    path = export_xlsx(rows, tmp_path / "out.xlsx")
    assert path.exists()
    from openpyxl import load_workbook

    wb = load_workbook(path)
    headers = [c.value for c in next(wb.active.iter_rows(min_row=1, max_row=1))]
    assert headers == EXCEL_HEADERS


def test_json_export_roundtrip(tmp_path: Path):
    product = Product(name="P", source_link="https://shopee.co.id/x-i.1.2", price=10)
    rows = flatten_product(product)
    path = export_json(rows, tmp_path / "out.json")
    import json

    data = json.loads(path.read_text(encoding="utf-8"))
    assert list(data[0].keys()) == EXCEL_HEADERS
